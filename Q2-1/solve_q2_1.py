"""2026 数模 C 题问题 2 改进方案核心代码。

改进点：
1. 仅用 1 月 18—31 日选择净负荷分位数和价格释放强度，避免使用正式评价期数据调参；
2. 低价时保留电池，高价时允许突破原计划 SOC 备用线，优先减少高价紧急购电；
3. 输出题目 result2 模板，并与 Q2 现行结果生成一个简明 CSV 对比表。
"""

from collections import defaultdict
from csv import writer
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import coo_matrix


PROJECT_DIR = Path(__file__).resolve().parent.parent
INPUT1_FILE = PROJECT_DIR / "附件1.xlsx"
INPUT2_FILE = PROJECT_DIR / "附件2.xlsx"
TEMPLATE_FILE = PROJECT_DIR / "附件5" / "result2.xlsx"
BASELINE_FILE = PROJECT_DIR / "Q2" / "result2.xlsx"
OUTPUT_FILE = Path(__file__).resolve().parent / "result2.xlsx"
COMPARE_FILE = Path(__file__).resolve().parent / "comparison.csv"

DT = 1 / 6
ETA = 0.90
POWER_MAX = 5000
INTERVAL_ENERGY_MAX = POWER_MAX * DT
ENERGY_MIN = 1200
ENERGY_MAX = 10800
ENERGY_INITIAL = 6000

START_DAY = 31
RESIDUAL_WINDOW = 56
QUANTILE_CANDIDATES = [0.70, 0.75, 0.80, 0.85, 0.90]
GAMMA_CANDIDATES = [0.20, 0.25, 0.30, 0.35, 0.50]
ZERO_TOL = 1e-6


# 读取电价和全年负载、光伏。
price_book = load_workbook(INPUT1_FILE, data_only=True, read_only=True)
price = np.array(
    [row[1] for row in price_book.active.iter_rows(min_row=2, values_only=True)],
    dtype=float,
)
price_book.close()

data_book = load_workbook(INPUT2_FILE, data_only=True, read_only=True)
load_rows = list(data_book["小区负载"].iter_rows(min_row=2, values_only=True))
pv_rows = list(
    data_book["光伏发电实际功率"].iter_rows(min_row=2, values_only=True)
)
dates = [row[0] for row in load_rows]
pv_dates = [row[0] for row in pv_rows]
load_power = np.array([row[1:] for row in load_rows], dtype=float)
pv_power = np.array([row[1:] for row in pv_rows], dtype=float)
data_book.close()

if price.shape != (144,) or np.any(~np.isfinite(price)) or np.any(price <= 0):
    raise ValueError("附件1中的电价应为144个正数。")
if load_power.shape != (365, 144) or pv_power.shape != (365, 144):
    raise ValueError("附件2的负载和光伏都应为365×144。")
if dates != pv_dates:
    raise ValueError("附件2的负载日期与光伏日期不一致。")
if np.any(~np.isfinite(load_power)) or np.any(~np.isfinite(pv_power)):
    raise ValueError("附件2中存在空值或非数值。")
if np.any(load_power < 0) or np.any(pv_power < 0):
    raise ValueError("附件2中出现负功率。")


# 48小时计划购电 LP 的稀疏矩阵固定不变，只更新每日右端项和初始 SOC。
N = 288
index = np.arange(N)

A_ub = coo_matrix(
    (
        np.concatenate([-np.ones(N), np.ones(N), -np.ones(N)]),
        (
            np.tile(index, 3),
            np.concatenate([index, N + index, 2 * N + index]),
        ),
    ),
    shape=(N, 4 * N),
).tocsr()

A_eq = coo_matrix(
    (
        np.concatenate(
            [
                -ETA * np.ones(N),
                np.ones(N) / ETA,
                np.ones(N),
                -np.ones(N - 1),
            ]
        ),
        (
            np.concatenate([index, index, index, index[1:]]),
            np.concatenate(
                [N + index, 2 * N + index, 3 * N + index, 3 * N + index[:-1]]
            ),
        ),
    ),
    shape=(N, 4 * N),
).tocsr()

objective = np.zeros(4 * N)
objective[:N] = np.tile(price, 2)
bounds = (
    [(0, None)] * N
    + [(0, INTERVAL_ENERGY_MAX)] * N
    + [(0, INTERVAL_ENERGY_MAX)] * N
    + [(ENERGY_MIN, ENERGY_MAX)] * (N - 1)
    + [(ENERGY_INITIAL, ENERGY_INITIAL)]
)


def solve_48h_plan(net_target, energy_start):
    """求48小时计划购电量及计划 SOC，只执行前24小时。"""
    b_eq = np.zeros(N)
    b_eq[0] = energy_start
    result = linprog(
        objective,
        A_ub=A_ub,
        b_ub=-net_target,
        A_eq=A_eq,
        b_eq=b_eq,
        bounds=bounds,
        method="highs-ds",
    )
    if not result.success:
        raise RuntimeError(f"日前线性规划求解失败：{result.message}")

    grid = result.x[:N].copy()
    charge = result.x[N:2 * N].copy()
    discharge = result.x[2 * N:3 * N].copy()

    # 消除光伏富余时可能出现的等价同时充放电解。
    both = (charge > 1e-8) & (discharge > 1e-8)
    amount = np.minimum(charge[both], discharge[both] / ETA**2)
    charge[both] -= amount
    discharge[both] -= ETA**2 * amount

    planned_soc = np.empty(N + 1)
    planned_soc[0] = energy_start
    for t in range(N):
        planned_soc[t + 1] = (
            planned_soc[t] + ETA * charge[t] - discharge[t] / ETA
        )

    if planned_soc.min() < ENERGY_MIN - 1e-5:
        raise RuntimeError("日前计划 SOC 低于下限。")
    if planned_soc.max() > ENERGY_MAX + 1e-5:
        raise RuntimeError("日前计划 SOC 高于上限。")
    return grid[:144], planned_soc[1:145]


# 严格滚动残差：每一天的误差只用于之后的日期。
load_error = np.full_like(load_power, np.nan)
pv_error = np.full_like(pv_power, np.nan)
for day in range(7, 365):
    load_error[day] = load_power[day] - load_power[day - 7]
    pv_error[day] = pv_power[day] - pv_power[day - 7:day].mean(axis=0)

price_score = (price - price.min()) / (price.max() - price.min())


def run_period(first_day, last_day, quantile, gamma, keep_detail=False):
    """按日期顺序运行，实际控制只读取当前及此前真实数据。"""
    day_count = last_day - first_day
    planned_grid = np.zeros((day_count, 144))
    actual_charge = np.zeros((day_count, 144))
    actual_discharge = np.zeros((day_count, 144))
    emergency_grid = np.zeros((day_count, 144))
    spill_grid = np.zeros((day_count, 144))
    soc_start = np.zeros(day_count)
    soc_end = np.zeros(day_count)
    daily_cost = np.zeros(day_count)

    current_soc = ENERGY_INITIAL
    for out_day, day in enumerate(range(first_day, last_day)):
        load_forecast_1 = load_power[day - 7]
        pv_forecast_1 = pv_power[day - 7:day].mean(axis=0)
        load_forecast_2 = load_power[day - 6]
        pv_forecast_2 = pv_forecast_1

        history = np.arange(max(7, day - RESIDUAL_WINDOW), day)
        if len(history) < 10:
            raise RuntimeError("可用历史残差日不足10天。")

        load_scenarios_1 = np.maximum(
            load_forecast_1[None, :] + load_error[history], 0
        )
        pv_scenarios_1 = np.maximum(
            pv_forecast_1[None, :] + pv_error[history], 0
        )
        load_scenarios_2 = np.maximum(
            load_forecast_2[None, :] + load_error[history], 0
        )
        pv_scenarios_2 = np.maximum(
            pv_forecast_2[None, :] + pv_error[history], 0
        )

        net_target = np.concatenate(
            [
                np.quantile(
                    load_scenarios_1 - pv_scenarios_1, quantile, axis=0
                ),
                np.quantile(
                    load_scenarios_2 - pv_scenarios_2, quantile, axis=0
                ),
            ]
        ) * DT

        grid, planned_soc = solve_48h_plan(net_target, current_soc)

        # 电价越高，备用线越低；低价时尽量留电，高价时优先放电。
        reserve_weight = (1 - price_score) ** gamma
        reserve = ENERGY_MIN + reserve_weight * (
            np.maximum(planned_soc, ENERGY_MIN) - ENERGY_MIN
        )

        charge = np.zeros(144)
        discharge = np.zeros(144)
        emergency = np.zeros(144)
        spill = np.zeros(144)
        actual_load = load_power[day] * DT
        actual_pv = pv_power[day] * DT

        soc = current_soc
        soc_start[out_day] = soc
        for t in range(144):
            balance = grid[t] + actual_pv[t] - actual_load[t]
            if balance >= 0:
                charge[t] = min(
                    balance,
                    INTERVAL_ENERGY_MAX,
                    max((ENERGY_MAX - soc) / ETA, 0),
                )
                spill[t] = balance - charge[t]
                soc += ETA * charge[t]
            else:
                deficit = -balance
                available = max(ETA * (soc - reserve[t]), 0)
                discharge[t] = min(deficit, INTERVAL_ENERGY_MAX, available)
                emergency[t] = deficit - discharge[t]
                soc -= discharge[t] / ETA

            if soc < ENERGY_MIN - 1e-5 or soc > ENERGY_MAX + 1e-5:
                raise RuntimeError(
                    f"{dates[day]:%Y-%m-%d} 第{t + 1}区间 SOC 越界：{soc:.6f}。"
                )

        planned_grid[out_day] = grid
        actual_charge[out_day] = charge
        actual_discharge[out_day] = discharge
        emergency_grid[out_day] = emergency
        spill_grid[out_day] = spill
        soc_end[out_day] = soc
        daily_cost[out_day] = price @ grid + 5 * price @ emergency
        current_soc = soc

        if keep_detail and ((out_day + 1) % 30 == 0 or day == last_day - 1):
            print(f"已完成 {out_day + 1}/{day_count} 天")

    return {
        "planned_grid": planned_grid,
        "charge": actual_charge,
        "discharge": actual_discharge,
        "emergency": emergency_grid,
        "spill": spill_grid,
        "soc_start": soc_start,
        "soc_end": soc_end,
        "daily_cost": daily_cost,
    }


# 只用正式评价期之前的1月18—31日选择两个参数。
best = None
for candidate_quantile in QUANTILE_CANDIDATES:
    for candidate_gamma in GAMMA_CANDIDATES:
        calibration = run_period(
            17, 31, candidate_quantile, candidate_gamma, keep_detail=False
        )
        calibration_cost = float(calibration["daily_cost"].sum())
        candidate = (calibration_cost, candidate_quantile, candidate_gamma)
        if best is None or candidate < best:
            best = candidate

_, selected_quantile, selected_gamma = best
print(
    f"1月校准完成：分位数={selected_quantile:.2f}，"
    f"价格释放参数={selected_gamma:.2f}"
)

result = run_period(
    START_DAY, 365, selected_quantile, selected_gamma, keep_detail=True
)
output_dates = dates[START_DAY:]
planned_grid = result["planned_grid"]
actual_charge = result["charge"]
actual_discharge = result["discharge"]
emergency_grid = result["emergency"]
spill_grid = result["spill"]
soc_start = result["soc_start"]
soc_end = result["soc_end"]
daily_cost = result["daily_cost"]


def merge_emergency_intervals(emergency, interval_labels):
    """把连续的十分钟紧急购电合并成一个时间段。"""
    groups = []
    start = None
    for t in range(len(emergency) + 1):
        active = t < len(emergency) and emergency[t] > ZERO_TOL
        if active and start is None:
            start = t
        if not active and start is not None:
            end = t - 1
            left = interval_labels[start].split("-", 1)[0]
            right = interval_labels[end].split("-", 1)[1]
            groups.append((f"{left}-{right}", float(emergency[start:t].sum())))
            start = None
    return groups


# 写入题目给定的 result2 模板。
output_book = load_workbook(TEMPLATE_FILE)
purchase_sheet = output_book["计划购电量"]
storage_sheet = output_book["充放电量"]
emergency_sheet = output_book["紧急购电量"]
interval_labels = [purchase_sheet.cell(1, col).value for col in range(2, 146)]

for i in range(len(output_dates)):
    row = i + 2
    if purchase_sheet.cell(row, 1).value.date() != output_dates[i].date():
        raise ValueError(f"result2模板第{row}行日期与附件2不一致。")
    for t in range(144):
        purchase_sheet.cell(row, t + 2, round(float(planned_grid[i, t]), 6))
    purchase_sheet.cell(
        row,
        146,
        round(float(planned_grid[i].sum() + emergency_grid[i].sum()), 6),
    )
    purchase_sheet.cell(row, 147, round(float(daily_cost[i]), 6))

if storage_sheet.max_row > 1:
    storage_sheet.delete_rows(2, storage_sheet.max_row - 1)
if emergency_sheet.max_row > 1:
    emergency_sheet.delete_rows(2, emergency_sheet.max_row - 1)

block_labels = [
    "0:00-4:00",
    "4:00-8:00",
    "8:00-12:00",
    "12:00-16:00",
    "16:00-20:00",
    "20:00-24:00",
]

row = 2
for i, date_value in enumerate(output_dates):
    for block in range(6):
        start = block * 24
        end = start + 24
        storage_sheet.cell(row, 1, date_value if block == 0 else None)
        storage_sheet.cell(row, 2, block_labels[block])
        storage_sheet.cell(row, 3, round(float(actual_charge[i, start:end].sum()), 6))
        storage_sheet.cell(
            row, 4, round(float(actual_discharge[i, start:end].sum()), 6)
        )
        if block == 0:
            storage_sheet.cell(row, 5, "0:00")
            storage_sheet.cell(row, 6, round(float(soc_start[i]), 6))
        elif block == 1:
            storage_sheet.cell(row, 5, "24:00")
            storage_sheet.cell(row, 6, round(float(soc_end[i]), 6))
        row += 1

row = 2
for i, date_value in enumerate(output_dates):
    groups = merge_emergency_intervals(emergency_grid[i], interval_labels)
    for group_index, (time_range, amount) in enumerate(groups):
        emergency_sheet.cell(row, 1, date_value if group_index == 0 else None)
        emergency_sheet.cell(row, 2, time_range)
        emergency_sheet.cell(row, 3, round(amount, 6))
        row += 1

output_book.save(OUTPUT_FILE)


# 读取现行结果，计算同口径指标。
baseline_book = load_workbook(BASELINE_FILE, data_only=True, read_only=True)
baseline_purchase, baseline_storage, baseline_emergency = baseline_book.worksheets[:3]
baseline_q = np.array(
    [row[1:145] for row in baseline_purchase.iter_rows(min_row=2, values_only=True)],
    dtype=float,
)
baseline_total_cost = sum(
    float(row[146])
    for row in baseline_purchase.iter_rows(min_row=2, values_only=True)
)

baseline_emergency_energy = 0.0
baseline_emergency_days = set()
current_date = None
for row in baseline_emergency.iter_rows(min_row=2, values_only=True):
    if row[0] is not None:
        current_date = row[0]
    amount = float(row[2] or 0)
    baseline_emergency_energy += amount
    if amount > ZERO_TOL:
        baseline_emergency_days.add(current_date)

baseline_charge = 0.0
baseline_discharge = 0.0
baseline_final_soc = None
for row in baseline_storage.iter_rows(min_row=2, values_only=True):
    baseline_charge += float(row[2] or 0)
    baseline_discharge += float(row[3] or 0)
    if row[4] == "24:00":
        baseline_final_soc = float(row[5])
baseline_book.close()

baseline_plan_cost = float((baseline_q * price).sum())
baseline_emergency_cost = baseline_total_cost - baseline_plan_cost
total_load = float(load_power[START_DAY:].sum() * DT)
total_pv = float(pv_power[START_DAY:].sum() * DT)
baseline_spill = (
    float(baseline_q.sum())
    + baseline_emergency_energy
    + total_pv
    + baseline_discharge
    - total_load
    - baseline_charge
)

improved_plan_energy = float(planned_grid.sum())
improved_emergency_energy = float(emergency_grid.sum())
improved_plan_cost = float((planned_grid * price).sum())
improved_emergency_cost = float((5 * emergency_grid * price).sum())
improved_total_cost = improved_plan_cost + improved_emergency_cost
improved_spill = float(spill_grid.sum())
improved_emergency_days = int(
    np.sum(emergency_grid.sum(axis=1) > ZERO_TOL)
)
saved_cost = baseline_total_cost - improved_total_cost

with COMPARE_FILE.open("w", newline="", encoding="utf-8-sig") as file:
    csv_writer = writer(file)
    csv_writer.writerow(
        [
            "方案",
            "分位数",
            "价格释放参数",
            "计划购电量(kWh)",
            "紧急购电量(kWh)",
            "计划购电费(元)",
            "紧急购电费(元)",
            "总购电费(元)",
            "弃电量(kWh)",
            "紧急购电天数",
            "年末储电量(kWh)",
            "相对现行节省(元)",
        ]
    )
    csv_writer.writerow(
        [
            "现行Q2",
            0.80,
            "计划SOC硬备用线",
            float(baseline_q.sum()),
            baseline_emergency_energy,
            baseline_plan_cost,
            baseline_emergency_cost,
            baseline_total_cost,
            baseline_spill,
            len(baseline_emergency_days),
            baseline_final_soc,
            0.0,
        ]
    )
    csv_writer.writerow(
        [
            "改进Q2-1",
            selected_quantile,
            selected_gamma,
            improved_plan_energy,
            improved_emergency_energy,
            improved_plan_cost,
            improved_emergency_cost,
            improved_total_cost,
            improved_spill,
            improved_emergency_days,
            float(soc_end[-1]),
            saved_cost,
        ]
    )

print("求解完成")
print(f"采用分位数：{selected_quantile:.2f}")
print(f"价格释放参数：{selected_gamma:.2f}")
print(f"计划购电量：{improved_plan_energy:.6f} kWh")
print(f"紧急购电量：{improved_emergency_energy:.6f} kWh")
print(f"计划购电费：{improved_plan_cost:.6f} 元")
print(f"紧急购电费：{improved_emergency_cost:.6f} 元")
print(f"总购电费：{improved_total_cost:.6f} 元")
print(f"估算弃电量：{improved_spill:.6f} kWh")
print(f"紧急购电天数：{improved_emergency_days} 天")
print(f"年末储电量：{soc_end[-1]:.6f} kWh")
print(f"比现行方案节省：{saved_cost:.6f} 元")
print(f"结果文件：{OUTPUT_FILE}")
print(f"对比文件：{COMPARE_FILE}")
