"""2026 数模 C 题问题 2 核心代码。

采用滚动因果预测、历史整日联合残差的 80% 净负荷分位数、
48 小时确定性线性规划，以及按真实负载/光伏逐时执行的储能反馈。
"""

from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog


PROJECT_DIR = Path(__file__).resolve().parent.parent
INPUT1_FILE = PROJECT_DIR / "附件1.xlsx"
INPUT2_FILE = PROJECT_DIR / "附件2.xlsx"
TEMPLATE_FILE = PROJECT_DIR / "附件5" / "result2.xlsx"
OUTPUT_FILE = Path(__file__).resolve().parent / "result2.xlsx"

DT = 1 / 6
ETA = 0.90
POWER_MAX = 5000
INTERVAL_ENERGY_MAX = POWER_MAX * DT
ENERGY_MIN = 1200
ENERGY_MAX = 10800
ENERGY_INITIAL = 6000

START_DAY = 31                 # 2025-02-01 在全年数据中的下标
QUANTILE = 0.80                # 紧急电价为普通电价5倍，对应单时段临界分位数
RESIDUAL_WINDOW = 56           # 最多使用此前56个完整日的联合残差
ZERO_TOL = 1e-6


def solve_48h_plan(net_target, price_48h, energy_start):
    """对未来48小时的分位数净负荷求LP，只执行前24小时计划。"""
    n = len(net_target)
    q0, c0, d0, e0 = 0, n, 2 * n, 3 * n
    variable_count = 4 * n

    objective = np.zeros(variable_count)
    objective[q0:q0 + n] = price_48h

    # q + d >= 净负荷 + c，即 c - d - q <= -净负荷。
    A_ub = np.zeros((n, variable_count))
    b_ub = -net_target
    for t in range(n):
        A_ub[t, q0 + t] = -1
        A_ub[t, c0 + t] = 1
        A_ub[t, d0 + t] = -1

    # E[t] = E[t-1] + eta*c[t] - d[t]/eta。
    A_eq = np.zeros((n, variable_count))
    b_eq = np.zeros(n)
    for t in range(n):
        A_eq[t, c0 + t] = -ETA
        A_eq[t, d0 + t] = 1 / ETA
        A_eq[t, e0 + t] = 1
        if t == 0:
            b_eq[t] = energy_start
        else:
            A_eq[t, e0 + t - 1] = -1

    # 不强制每天闭合；仅令48小时前瞻末端回到名义电量，避免免费耗尽电池。
    bounds = (
        [(0, None)] * n
        + [(0, INTERVAL_ENERGY_MAX)] * n
        + [(0, INTERVAL_ENERGY_MAX)] * n
        + [(ENERGY_MIN, ENERGY_MAX)] * (n - 1)
        + [(ENERGY_INITIAL, ENERGY_INITIAL)]
    )

    result = linprog(
        objective,
        A_ub=A_ub,
        b_ub=b_ub,
        A_eq=A_eq,
        b_eq=b_eq,
        bounds=bounds,
        method="highs-ds",
    )
    if not result.success:
        raise RuntimeError(f"日前线性规划求解失败：{result.message}")

    grid = result.x[q0:q0 + n].copy()
    charge = result.x[c0:c0 + n].copy()
    discharge = result.x[d0:d0 + n].copy()

    # 消除LP在光伏富余时可能出现的等价同时充放电解。
    for t in range(n):
        if charge[t] > 1e-8 and discharge[t] > 1e-8:
            amount = min(charge[t], discharge[t] / ETA**2)
            charge[t] -= amount
            discharge[t] -= ETA**2 * amount

    grid[np.abs(grid) < 1e-8] = 0
    charge[np.abs(charge) < 1e-8] = 0
    discharge[np.abs(discharge) < 1e-8] = 0

    planned_soc = np.empty(n + 1)
    planned_soc[0] = energy_start
    for t in range(n):
        planned_soc[t + 1] = (
            planned_soc[t] + ETA * charge[t] - discharge[t] / ETA
        )

    margin = grid + discharge - net_target - charge
    if margin.min() < -1e-5:
        raise RuntimeError("日前计划供需约束校验失败。")
    if planned_soc.min() < ENERGY_MIN - 1e-5 or planned_soc.max() > ENERGY_MAX + 1e-5:
        raise RuntimeError("日前计划SOC上下界校验失败。")
    if abs(planned_soc[-1] - ENERGY_INITIAL) > 1e-5:
        raise RuntimeError("48小时前瞻末端SOC校验失败。")

    return grid, planned_soc


def merge_emergency_intervals(emergency, interval_labels):
    """把连续的十分钟紧急购电合并成一个时间段。"""
    groups = []
    start = None
    for t in range(len(emergency) + 1):
        active = t < len(emergency) and emergency[t] > ZERO_TOL
        if active and start is None:
            start = t
        if not active and start is not None:
            end = t - 1
            left = interval_labels[start].split("-", 1)[0]
            right = interval_labels[end].split("-", 1)[1]
            groups.append((f"{left}-{right}", float(emergency[start:t].sum())))
            start = None
    return groups


# 读取固定的分时电价。
price_book = load_workbook(INPUT1_FILE, data_only=True, read_only=True)
price_rows = list(price_book.active.iter_rows(min_row=2, values_only=True))
price = np.array([row[1] for row in price_rows], dtype=float)
price_book.close()
if price.shape != (144,) or np.any(~np.isfinite(price)) or np.any(price <= 0):
    raise ValueError("附件1中的电价应为144个正数。")

# 读取2025年逐日负载和光伏。附件列与结果模板按原有顺序一一对应。
data_book = load_workbook(INPUT2_FILE, data_only=True, read_only=True)
load_sheet = data_book["小区负载"]
pv_sheet = data_book["光伏发电实际功率"]

load_rows = list(load_sheet.iter_rows(min_row=2, values_only=True))
pv_rows = list(pv_sheet.iter_rows(min_row=2, values_only=True))
dates = [row[0] for row in load_rows]
pv_dates = [row[0] for row in pv_rows]
load_power = np.array([row[1:] for row in load_rows], dtype=float)
pv_power = np.array([row[1:] for row in pv_rows], dtype=float)
data_book.close()

if load_power.shape != (365, 144) or pv_power.shape != (365, 144):
    raise ValueError(
        f"附件2数据形状错误：负载{load_power.shape}，光伏{pv_power.shape}。"
    )
if dates != pv_dates:
    raise ValueError("附件2的负载日期与光伏日期不一致。")
if np.any(~np.isfinite(load_power)) or np.any(~np.isfinite(pv_power)):
    raise ValueError("附件2中存在空值或非数值。")
if np.any(load_power < 0) or np.any(pv_power < 0):
    raise ValueError("附件2中出现负功率。")

# 预先计算严格滚动基线残差。每一天的残差只会在后续日期使用。
load_error = np.full_like(load_power, np.nan)
pv_error = np.full_like(pv_power, np.nan)
for day in range(7, 365):
    load_error[day] = load_power[day] - load_power[day - 7]
    pv_error[day] = pv_power[day] - pv_power[day - 7:day].mean(axis=0)

output_dates = dates[START_DAY:]
day_count = len(output_dates)
planned_grid = np.zeros((day_count, 144))
actual_charge = np.zeros((day_count, 144))
actual_discharge = np.zeros((day_count, 144))
emergency_grid = np.zeros((day_count, 144))
soc_start = np.zeros(day_count)
soc_end = np.zeros(day_count)
daily_cost = np.zeros(day_count)

current_soc = ENERGY_INITIAL
price_48h = np.tile(price, 2)

for out_day, day in enumerate(range(START_DAY, 365)):
    # 当天负载用上周同刻；光伏用此前7个完整日的同刻均值。
    load_forecast_1 = load_power[day - 7]
    pv_forecast_1 = pv_power[day - 7:day].mean(axis=0)

    # 次日仅作为前瞻，不输出也不执行。其输入在当天0:00均已知或可预测。
    load_forecast_2 = load_power[day - 6]
    pv_forecast_2 = pv_forecast_1

    first_error_day = max(7, day - RESIDUAL_WINDOW)
    history = np.arange(first_error_day, day)
    if len(history) < 10:
        raise RuntimeError("可用历史残差日不足10天。")

    # 负载和光伏残差按整日成对加入，保留二者及日内时序相关性。
    load_scenarios_1 = np.maximum(
        load_forecast_1[None, :] + load_error[history], 0
    )
    pv_scenarios_1 = np.maximum(
        pv_forecast_1[None, :] + pv_error[history], 0
    )
    load_scenarios_2 = np.maximum(
        load_forecast_2[None, :] + load_error[history], 0
    )
    pv_scenarios_2 = np.maximum(
        pv_forecast_2[None, :] + pv_error[history], 0
    )

    net_target_1 = np.quantile(
        load_scenarios_1 - pv_scenarios_1, QUANTILE, axis=0
    ) * DT
    net_target_2 = np.quantile(
        load_scenarios_2 - pv_scenarios_2, QUANTILE, axis=0
    ) * DT
    net_target_48h = np.concatenate([net_target_1, net_target_2])

    grid_48h, planned_soc_48h = solve_48h_plan(
        net_target_48h, price_48h, current_soc
    )
    grid = grid_48h[:144]

    # 计划SOC轨迹作为因果备用线。实际控制只使用当前缺口和当前SOC，
    # 不读取当天未来真实负载或光伏。
    reserve = np.maximum(planned_soc_48h[1:145], ENERGY_MIN)
    charge = np.zeros(144)
    discharge = np.zeros(144)
    emergency = np.zeros(144)

    soc = current_soc
    soc_start[out_day] = soc
    actual_load = load_power[day] * DT
    actual_pv = pv_power[day] * DT

    for t in range(144):
        balance = grid[t] + actual_pv[t] - actual_load[t]
        if balance >= 0:
            charge[t] = min(
                balance,
                INTERVAL_ENERGY_MAX,
                max((ENERGY_MAX - soc) / ETA, 0),
            )
            soc += ETA * charge[t]
        else:
            deficit = -balance
            available = max(ETA * (soc - reserve[t]), 0)
            discharge[t] = min(deficit, INTERVAL_ENERGY_MAX, available)
            emergency[t] = deficit - discharge[t]
            soc -= discharge[t] / ETA

        if soc < ENERGY_MIN - 1e-5 or soc > ENERGY_MAX + 1e-5:
            raise RuntimeError(
                f"{dates[day]:%Y-%m-%d} 第{t + 1}区间SOC越界：{soc:.6f}。"
            )

    planned_grid[out_day] = grid
    actual_charge[out_day] = charge
    actual_discharge[out_day] = discharge
    emergency_grid[out_day] = emergency
    soc_end[out_day] = soc
    daily_cost[out_day] = price @ grid + 5 * price @ emergency
    current_soc = soc

    if (out_day + 1) % 30 == 0 or out_day == day_count - 1:
        print(f"已完成 {out_day + 1}/{day_count} 天")

# 写入题目给定模板的副本。
output_book = load_workbook(TEMPLATE_FILE)
purchase_sheet = output_book["计划购电量"]
storage_sheet = output_book["充放电量"]
emergency_sheet = output_book["紧急购电量"]

interval_labels = [purchase_sheet.cell(1, col).value for col in range(2, 146)]
if len(interval_labels) != 144 or any(label is None for label in interval_labels):
    raise ValueError("result2模板的144个时间段表头不完整。")

for i in range(day_count):
    row = i + 2
    template_date = purchase_sheet.cell(row, 1).value
    if template_date.date() != output_dates[i].date():
        raise ValueError(f"result2模板第{row}行日期与附件2不一致。")
    for t in range(144):
        purchase_sheet.cell(row, t + 2, round(float(planned_grid[i, t]), 6))

    # 全天购电量包含计划购电和紧急购电；全天购电费按题意计入5倍紧急电价。
    purchase_sheet.cell(
        row, 146,
        round(float(planned_grid[i].sum() + emergency_grid[i].sum()), 6),
    )
    purchase_sheet.cell(row, 147, round(float(daily_cost[i]), 6))

# 模板后两张工作表只有示意行，按全年真实结果重新填充。
if storage_sheet.max_row > 1:
    storage_sheet.delete_rows(2, storage_sheet.max_row - 1)

if emergency_sheet.max_row > 1:
    emergency_sheet.delete_rows(2, emergency_sheet.max_row - 1)

block_labels = [
    "0:00-4:00", "4:00-8:00", "8:00-12:00",
    "12:00-16:00", "16:00-20:00", "20:00-24:00",
]

row = 2
for i, date_value in enumerate(output_dates):
    for block in range(6):
        start = block * 24
        end = start + 24
        storage_sheet.cell(row, 1, date_value if block == 0 else None)
        storage_sheet.cell(row, 2, block_labels[block])
        storage_sheet.cell(row, 3, round(float(actual_charge[i, start:end].sum()), 6))
        storage_sheet.cell(row, 4, round(float(actual_discharge[i, start:end].sum()), 6))
        if block == 0:
            storage_sheet.cell(row, 5, "0:00")
            storage_sheet.cell(row, 6, round(float(soc_start[i]), 6))
        elif block == 1:
            storage_sheet.cell(row, 5, "24:00")
            storage_sheet.cell(row, 6, round(float(soc_end[i]), 6))
        row += 1

row = 2
for i, date_value in enumerate(output_dates):
    groups = merge_emergency_intervals(emergency_grid[i], interval_labels)
    for group_index, (time_range, amount) in enumerate(groups):
        emergency_sheet.cell(row, 1, date_value if group_index == 0 else None)
        emergency_sheet.cell(row, 2, time_range)
        emergency_sheet.cell(row, 3, round(amount, 6))
        row += 1

output_book.save(OUTPUT_FILE)

normal_energy = planned_grid.sum()
emergency_energy = emergency_grid.sum()
normal_cost = np.tile(price, (day_count, 1))
normal_cost = float((normal_cost * planned_grid).sum())
emergency_cost = float((5 * np.tile(price, (day_count, 1)) * emergency_grid).sum())

print("求解完成")
print(f"计划购电量：{normal_energy:.6f} kWh")
print(f"紧急购电量：{emergency_energy:.6f} kWh")
print(f"计划购电费：{normal_cost:.6f} 元")
print(f"紧急购电费：{emergency_cost:.6f} 元")
print(f"总购电费：{normal_cost + emergency_cost:.6f} 元")
print(f"年末储电量：{current_soc:.6f} kWh")
print(f"结果文件：{OUTPUT_FILE}")
