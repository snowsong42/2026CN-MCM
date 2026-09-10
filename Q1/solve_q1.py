"""2026 数模 C 题问题 1：确定性线性规划。"""

from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog


# 路径
PROJECT_DIR = Path(__file__).resolve().parent.parent
INPUT_FILE = PROJECT_DIR / "附件1.xlsx"
TEMPLATE_FILE = PROJECT_DIR / "附件5" / "result1.xlsx"
OUTPUT_FILE = Path(__file__).resolve().parent / "result1.xlsx"

# 储能参数
DT = 1 / 6                  # 10 分钟，单位：小时
ETA = 0.90                  # 充电效率和放电效率均为 90%
POWER_MAX = 5000            # kW
ENERGY_MIN = 1200           # kWh
ENERGY_MAX = 10800          # kWh
ENERGY_INITIAL = 6000       # kWh


# 读取附件1。其 144 行数据与结果模板的 144 行按原有顺序一一对应。
input_book = load_workbook(INPUT_FILE, data_only=True, read_only=True)
input_sheet = input_book.active
rows = list(input_sheet.iter_rows(min_row=2, values_only=True))

if len(rows) != 144:
    raise ValueError(f"附件1应包含144个十分钟数据点，实际读取到{len(rows)}个。")

price = np.array([row[1] for row in rows], dtype=float)
load_energy = np.array([row[2] for row in rows], dtype=float) * DT
pv_energy = np.array([row[3] for row in rows], dtype=float) * DT

if np.any(~np.isfinite(price)) or np.any(~np.isfinite(load_energy)) or np.any(~np.isfinite(pv_energy)):
    raise ValueError("附件1中存在空值或非数值。")

n = len(rows)
q0 = 0
c0 = n
d0 = 2 * n
e0 = 3 * n
variable_count = 4 * n

# 变量依次为：购电量 q、充电量 c、放电量 d、区间末储电量 E。
# 目标：min sum(price[t] * q[t])。
objective = np.zeros(variable_count)
objective[q0:q0 + n] = price

# 供需约束：q + 光伏 + d >= 负载 + c。
# 多余电能允许不利用，因此写成不等式，不额外设置弃电变量。
A_ub = np.zeros((n, variable_count))
b_ub = pv_energy - load_energy
for t in range(n):
    A_ub[t, q0 + t] = -1
    A_ub[t, c0 + t] = 1
    A_ub[t, d0 + t] = -1

# SOC 状态方程：E[t] = E[t-1] + eta*c[t] - d[t]/eta。
A_eq = np.zeros((n, variable_count))
b_eq = np.zeros(n)
for t in range(n):
    A_eq[t, c0 + t] = -ETA
    A_eq[t, d0 + t] = 1 / ETA
    A_eq[t, e0 + t] = 1
    if t == 0:
        b_eq[t] = ENERGY_INITIAL
    else:
        A_eq[t, e0 + t - 1] = -1

interval_energy_max = POWER_MAX * DT
bounds = (
    [(0, None)] * n
    + [(0, interval_energy_max)] * n
    + [(0, interval_energy_max)] * n
    + [(ENERGY_MIN, ENERGY_MAX)] * (n - 1)
    + [(ENERGY_INITIAL, ENERGY_INITIAL)]
)

result = linprog(
    objective,
    A_ub=A_ub,
    b_ub=b_ub,
    A_eq=A_eq,
    b_eq=b_eq,
    bounds=bounds,
    method="highs-ds",
)

if not result.success:
    raise RuntimeError(f"线性规划求解失败：{result.message}")

grid = result.x[q0:q0 + n].copy()
charge = result.x[c0:c0 + n].copy()
discharge = result.x[d0:d0 + n].copy()

# LP 可能在光伏富余时给出等价的同时充放电解。按论文中的等价变换消除它，
# 不改变 SOC 和最优费用，只增加未利用的剩余电能。
for t in range(n):
    if charge[t] > 1e-8 and discharge[t] > 1e-8:
        amount = min(charge[t], discharge[t] / ETA**2)
        charge[t] -= amount
        discharge[t] -= ETA**2 * amount

grid[np.abs(grid) < 1e-8] = 0
charge[np.abs(charge) < 1e-8] = 0
discharge[np.abs(discharge) < 1e-8] = 0

# 独立复算 SOC 并检查物理约束。
soc = np.empty(n + 1)
soc[0] = ENERGY_INITIAL
for t in range(n):
    soc[t + 1] = soc[t] + ETA * charge[t] - discharge[t] / ETA

supply_margin = grid + pv_energy + discharge - load_energy - charge
if supply_margin.min() < -1e-5:
    raise RuntimeError(f"供需约束校验失败，最小余量为{supply_margin.min():.6f} kWh。")
if soc.min() < ENERGY_MIN - 1e-5 or soc.max() > ENERGY_MAX + 1e-5:
    raise RuntimeError("储电量上下界校验失败。")
if abs(soc[-1] - ENERGY_INITIAL) > 1e-5:
    raise RuntimeError(f"24:00储电量未回到初值，偏差为{soc[-1] - ENERGY_INITIAL:.6f} kWh。")

# 写入题目给定模板，不改动原始附件和模板文件。
output_book = load_workbook(TEMPLATE_FILE)
purchase_sheet = output_book["计划购电量"]
storage_sheet = output_book["充放电量"]

for t in range(n):
    purchase_sheet.cell(row=t + 2, column=2, value=round(float(grid[t]), 6))

# 每 24 个十分钟区间合计为 4 小时。
for block in range(6):
    start = block * 24
    end = start + 24
    storage_sheet.cell(row=block + 2, column=2, value=round(float(charge[start:end].sum()), 6))
    storage_sheet.cell(row=block + 2, column=3, value=round(float(discharge[start:end].sum()), 6))

storage_sheet.cell(row=2, column=5, value=round(float(soc[0]), 6))
storage_sheet.cell(row=3, column=5, value=round(float(soc[-1]), 6))
output_book.save(OUTPUT_FILE)

total_grid = grid.sum()
total_cost = price @ grid
print(f"求解成功：{result.message}")
print(f"全天购电量：{total_grid:.6f} kWh")
print(f"全天购电费：{total_cost:.6f} 元")
print(f"最低/最高储电量：{soc.min():.6f} / {soc.max():.6f} kWh")
print(f"结果文件：{OUTPUT_FILE}")
